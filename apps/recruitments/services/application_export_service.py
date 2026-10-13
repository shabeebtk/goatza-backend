# recruitments/services/application_export_service.py
"""
The applicants spreadsheet an org downloads — one .xlsx, at most 100 rows.

THIS FILE HOLDS MINORS' CONTACT DETAILS. Who may build one (owner/admin) is
the view's job, and the audit line it writes is the only trail; this module
only turns rows the caller already chose into bytes, and logs nothing.

FORMULA INJECTION. Names, answers and phone numbers are typed by players. A
spreadsheet treats a cell that starts with ``=`` (or ``+ - @``) as a formula,
and openpyxl turns any such string into one on assignment — so a player called
``=HYPERLINK("http://…","Click")`` would hand the org a live link. Every
player-supplied value goes through ``_write_text``, which pins the cell to the
string type and sets Excel's own quote-prefix flag, so it is stored and shown
as the text it is.
"""

import io
import re

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.recruitments.legacy_status import LEGACY_STATUSES
from apps.recruitments.models import Recruitment, RecruitmentApplication
from apps.recruitments.serializers.recruitment_list_serializers import (
    application_session_payload,
)
from utils.transactional_emails import (
    STATUS_EMAIL_COPY,
    format_clock,
    format_date,
)

# One file never holds more than this — the founder's rule, and what keeps a
# leaked file's blast radius small.
MAX_ROWS = 100

SHEET_TITLE = "Applicants"

# What a spreadsheet reads as the start of a formula. Tab and CR are on the
# OWASP list too: some importers strip them and then see the "=" behind.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

# Column widths, in characters. Capped so one long answer does not push the
# rest of the sheet off screen; the full text is still in the cell.
_MIN_WIDTH = 8
_MAX_WIDTH = 50

_STATUS_LABELS = dict(RecruitmentApplication.Status.choices)
_SELF_OUTCOME_LABELS = dict(RecruitmentApplication.SelfOutcome.choices)

# The app's own rule for a category's gender word (eligibility.ts,
# addedGenderWord): only when the CATEGORY names one, and not when its title
# already says it.
_TITLE_NAMES_GENDER = re.compile(r"(boys|girls|men|women)", re.IGNORECASE)


def _safe_text(value):
    """``value`` as plain text — never None, never a formula."""
    if value is None:
        return ""
    return str(value)


def _write_text(cell, value):
    """Write player text so a spreadsheet can only ever show it, not run it."""
    text = _safe_text(value)
    cell.value = text
    # openpyxl has already decided "=..." is a formula; this undoes that.
    cell.data_type = "s"
    if text.startswith(_FORMULA_PREFIXES):
        cell.quotePrefix = True


def _yes_no(value):
    if value is None:
        return ""
    return "Yes" if value else "No"


def _status_label(status):
    """The status as a person reads it, legacy values included."""
    if status in _STATUS_LABELS:
        return str(_STATUS_LABELS[status])
    if status in LEGACY_STATUSES:
        return STATUS_EMAIL_COPY[status]["badge_label"]
    return status.replace("_", " ").capitalize()


def _age_group_label(category):
    if category is None:
        return ""
    word = ""
    if category.gender in ("male", "female") and not _TITLE_NAMES_GENDER.search(
        category.title
    ):
        word = "Boys" if category.gender == "male" else "Girls"
    return " ".join(part for part in (category.title, word) if part)


def _session_label(application):
    """Place first, then the day, then the time — the list's centre chip."""
    payload = application_session_payload(application)
    if payload is None:
        return ""
    place = (payload["city"] or "").strip() or (payload["venue_name"] or "").strip()
    return " · ".join(
        part for part in (
            place,
            format_date(payload["date"]),
            format_clock(payload["start_time"]),
        ) if part
    )


def _applied_on(value, tzinfo):
    """ "12 Oct 2026, 4:05 PM" on the recruitment's own clock."""
    local = value.astimezone(tzinfo)
    hour = local.hour % 12 or 12
    meridiem = "AM" if local.hour < 12 else "PM"
    return f"{format_date(local)}, {hour}:{local.minute:02d} {meridiem}"


def _fee_marked_by(application):
    member = application.fee_marked_by
    if member is None:
        return ""
    profile = getattr(member.user, "profile", None)
    return getattr(profile, "name", "") or member.user.username or ""


def _answers_by_question(application):
    """{question_id: "text"} — checkbox options joined ", ", in their order."""
    values = {}
    for answer in application.answers.all():
        if answer.selected_option_id:
            text = answer.selected_option.value
        else:
            text = answer.answer_text
        if text:
            values.setdefault(answer.question_id, []).append(text)
    return {key: ", ".join(parts) for key, parts in values.items()}


def _columns(recruitment, questions):
    """
    ``[(header, getter, is_player_text)]`` for THIS recruitment — a column
    that cannot apply to it (no age groups, no fee, no centres to pick, no
    trial to report on) is left out rather than shipped empty.
    """
    frontend = (settings.FRONTEND_BASE_URL or "").rstrip("/")
    tz = recruitment.zoneinfo
    has_groups = recruitment.age_categories.exists()
    is_trial = recruitment.recruitment_type == Recruitment.Type.OPEN_TRIAL

    columns = [
        ("Name", lambda a: a.shared_name, True),
        ("Phone", lambda a: a.shared_phone, True),
        ("Email", lambda a: a.shared_email, True),
        ("Username", lambda a: a.applicant.username, True),
        ("Profile link", lambda a: f"{frontend}/{a.applicant.username}", True),
        ("Birth year", lambda a: getattr(a, "birth_year", None), False),
    ]
    if has_groups:
        columns.append(
            ("Age group", lambda a: _age_group_label(a.age_category), True)
        )
    if recruitment.session_mode == Recruitment.SessionMode.CHOOSE_ONE:
        columns.append(("Centre / date", _session_label, True))
    columns += [
        ("Status", lambda a: _status_label(a.status), False),
        ("Applied on", lambda a: _applied_on(a.applied_at, tz), False),
    ]
    if recruitment.is_paid:
        columns += [
            ("Fee paid", lambda a: _yes_no(a.fee_paid), False),
            ("Fee marked by", _fee_marked_by, True),
        ]
    if is_trial:
        columns += [
            ("Attended", lambda a: _yes_no(a.attended_self_reported), False),
            (
                "Outcome",
                lambda a: str(_SELF_OUTCOME_LABELS.get(a.outcome_self_reported, "")),
                False,
            ),
            ("Trial rating", lambda a: a.trial_rating, False),
        ]
    if has_groups:
        columns.append(
            ("Age mismatch", lambda a: _yes_no(a.age_mismatch_at_apply), False)
        )

    # One column per question, in the order the org asked them. The header
    # is the org's own wording, and the cell is the player's — both text.
    for question in questions:
        columns.append(
            (
                question.question,
                lambda a, qid=question.id: a._export_answers.get(qid, ""),
                True,
            )
        )

    return columns


def build_applicants_workbook(recruitment, applications, start):
    """
    The .xlsx for ``applications`` (already filtered, ordered and sliced by
    the caller — at most ``MAX_ROWS``), numbered from ``start``. Bytes.

    The caller prefetches answers → question + selected_option, so this is a
    handful of queries for 100 rows rather than one per answer.
    """
    questions = list(recruitment.questions.all())
    columns = _columns(recruitment, questions)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE

    headers = ["#"] + [header for header, _, _ in columns]
    widths = [len(header) for header in headers]

    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        # The org wrote the question text, so it is guarded like player text.
        _write_text(cell, header)
        cell.font = Font(bold=True)

    for offset, application in enumerate(applications[:MAX_ROWS]):
        application._export_answers = _answers_by_question(application)
        row = offset + 2

        sheet.cell(row=row, column=1, value=start + offset)

        for index, (header, getter, is_player_text) in enumerate(columns, start=2):
            value = getter(application)
            cell = sheet.cell(row=row, column=index)
            if is_player_text:
                _write_text(cell, value)
            else:
                cell.value = "" if value is None else value
            widths[index - 1] = max(widths[index - 1], len(_safe_text(value)))

        # A phone is an identifier, not a number: stored as text so a leading
        # + or 0 survives, and formatted as text so Excel leaves it alone.
        phone_cell = sheet.cell(row=row, column=headers.index("Phone") + 1)
        phone_cell.number_format = "@"

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            max(_MIN_WIDTH, min(width + 2, _MAX_WIDTH))
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
